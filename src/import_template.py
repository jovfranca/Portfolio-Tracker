"""Small fictional workbook for the transaction import download."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def transaction_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Transacoes'
    sheet.append([
        'ticker', 'broker', 'type', 'trade_date', 'settlement_date', 'quantity',
        'unit_price', 'transaction_currency', 'fx_rate', 'allocation_class',
        'brokerage_fee', 'other_fees', 'notes',
    ])
    sheet.append([
        'FICTICIO-BR', 'Corretora Fictícia', 'Buy', date(2024, 1, 2), date(2024, 1, 4),
        '10', '25.50', 'BRL', None, 'Exemplo', '1.25', '0.50',
        'Exemplo fictício: substitua ou remova antes de importar.',
    ])
    sheet.append([
        'FICTICIO-BR', 'Corretora Fictícia', 'Venda', date(2024, 2, 1), date(2024, 2, 5),
        '2', '30.00', 'BRL', None, 'Exemplo', '0', '0',
        'Exemplo fictício: venda parcial da compra acima.',
    ])
    sheet.append([
        'FICTICIO-US', 'Corretora Fictícia', 'Compra', date(2024, 3, 4), date(2024, 3, 6),
        '1.5', '100.25', 'USD', '5.25', 'Exemplo', None, None,
        'Exemplo fictício: FX informado, sem consulta externa.',
    ])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='285C46')
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        row[3].number_format = row[4].number_format = 'yyyy-mm-dd'
        sheet.row_dimensions[row[0].row].height = 42
    for number, width in enumerate([20, 25, 14, 18, 20, 16, 18, 20, 16, 22, 20, 18, 60], start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
