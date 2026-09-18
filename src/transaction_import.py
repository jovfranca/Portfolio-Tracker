"""Safe CSV/XLSX transaction parsing used by preview and confirmation."""
from csv import DictReader, Sniffer, Error as CSVError
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from pydantic import ValidationError
from sqlalchemy import select

from src.schemas import TransactionInput
from src.models import Transaction
from src.instruments import resolve_instrument
from src.services import transaction_values


MAX_IMPORT_BYTES = 5_000_000
MAX_IMPORT_ROWS = 5_000
MAX_XLSX_EXPANDED_BYTES = 25_000_000
ALIASES = {
    'ticker': 'asset',
    'transaction_type': 'type',
    'trade date': 'trade_date',
    'settlement date': 'settlement_date',
    'unit_price': 'price',
    'unit price': 'price',
    'currency': 'asset_currency',
    'fx rate': 'fx_rate',
    'asset currency': 'asset_currency',
    'allocation class': 'allocation_class',
    'brokerage fee': 'brokerage_fee',
    'other fees': 'other_fees',
}
ALLOWED_FIELDS = set(TransactionInput.model_fields)


def _header(value):
    name = str(value or '').strip().lower()
    return ALIASES.get(name, name)


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _csv_rows(content):
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError as error:
        raise ValueError('O CSV deve usar codificação UTF-8.') from error
    try:
        dialect = Sniffer().sniff(text[:4096], delimiters=',;\t')
    except Exception:
        dialect = 'excel'
    reader = DictReader(StringIO(text), dialect=dialect, strict=True)
    rows = []
    try:
        _validate_headers(reader.fieldnames or [])
        for row in reader:
            if None in row:
                raise ValueError(f'Linha {reader.line_num}: células sem coluna correspondente.')
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ValueError('O arquivo excede o limite de 5.000 transações.')
            rows.append((reader.line_num, row))
    except CSVError as error:
        raise ValueError(f'CSV inválido na linha {reader.line_num}: {error}') from error
    return rows


def _validate_headers(headers):
    names = [_header(value) for value in headers]
    if len(set(names)) != len(names):
        raise ValueError('O arquivo contém nomes de coluna duplicados ou ambíguos.')
    return names


def _xlsx_rows(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_XLSX_EXPANDED_BYTES:
                raise ValueError('A planilha descompactada excede o limite de 25 MB.')
    except BadZipFile as error:
        raise ValueError('Não foi possível ler a planilha XLSX.') from error
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError('Não foi possível ler a planilha XLSX.') from error
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError('A planilha não contém uma aba de dados.')
        # Ignore declared dimensions: a tiny archive can claim billions of cells.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            return []
        names = _validate_headers(headers)
        rows = []
        for number, values in enumerate(iterator, start=2):
            if number > MAX_IMPORT_ROWS + 1:
                raise ValueError('O arquivo excede o limite de 5.000 linhas.')
            if not any(value is not None for value in values):
                continue
            if any(value is not None for value in values[len(names):]):
                raise ValueError(f'Linha {number}: células sem coluna correspondente.')
            rows.append((number, dict(zip(names, values))))
        return rows
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Não foi possível ler a planilha XLSX.') from error
    finally:
        workbook.close()


def read_rows(filename, content, *, include_line_numbers=False):
    if not content:
        raise ValueError('O arquivo está vazio.')
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError('O arquivo excede o limite de 5 MB.')
    suffix = Path(filename).suffix.lower()
    if suffix == '.csv':
        raw_rows = _csv_rows(content)
    elif suffix == '.xlsx':
        raw_rows = _xlsx_rows(content)
    else:
        raise ValueError('Use um arquivo CSV ou XLSX.')
    if not raw_rows:
        raise ValueError('O arquivo não contém transações.')
    if len(raw_rows) > MAX_IMPORT_ROWS:
        raise ValueError('O arquivo excede o limite de 5.000 transações.')
    normalized = []
    for number, raw in raw_rows:
        row = {}
        for key, value in raw.items():
            name = _header(key)
            if name in ALLOWED_FIELDS and _cell(value) != '':
                row[name] = _cell(value)
        normalized.append((number, row) if include_line_numbers else row)
    return normalized


def _errors(error):
    return [
        {
            'field': '.'.join(str(item) for item in detail['loc']) or 'row',
            'message': detail['msg'],
        }
        for detail in error.errors(include_url=False)
    ]


def preview_import(session, filename, content, portfolio_id=None):
    result = []
    currencies = dict(session.execute(select(Transaction.instrument_id, Transaction.asset_currency).where(
        Transaction.portfolio_id == portfolio_id,
    )).all()) if portfolio_id is not None else {}
    for number, raw in read_rows(filename, content, include_line_numbers=True):
        try:
            # Imports must never infer BRL from an absent or misspelled column.
            raw.setdefault('asset_currency', '')
            payload = TransactionInput.model_validate(raw)
            resolution = None
            resolved_instrument_id = None
            if portfolio_id is not None:
                resolution = resolve_instrument(
                    session, payload.asset, currency=payload.asset_currency,
                )
                if resolution.status != 'resolved':
                    result.append({
                        'row': number, 'valid': False,
                        'data': payload.model_dump(mode='json'),
                        'instrument_resolution': resolution.status,
                        'errors': [{'field': 'asset', 'message':
                            f'Instrumento {resolution.status}; selecione uma correspondência explícita.'}],
                    })
                    continue
                resolved_instrument_id = resolution.instrument.id
                if resolution.instrument.currency != payload.asset_currency:
                    result.append({
                        'row': number, 'valid': False, 'instrument_resolution': 'resolved',
                        'errors': [{'field': 'asset_currency', 'message':
                            f'As cotações deste instrumento usam {resolution.instrument.currency}.'}],
                    })
                    continue
            existing = currencies.get(resolved_instrument_id)
            if existing is not None and existing != payload.asset_currency:
                result.append({
                    'row': number, 'valid': False,
                    'instrument_resolution': 'resolved',
                    'errors': [{'field': 'asset_currency', 'message':
                        f'O instrumento já está registrado em {existing}; não misture moedas na mesma posição.'}],
                })
                continue
            if resolved_instrument_id is not None:
                currencies[resolved_instrument_id] = payload.asset_currency
            values = transaction_values(session, payload)
            values.pop('date_time')
            normalized = TransactionInput.model_validate(values).model_dump(mode='json')
            if resolved_instrument_id is not None:
                normalized['instrument_id'] = resolved_instrument_id
        except ValidationError as error:
            result.append({'row': number, 'valid': False, 'errors': _errors(error)})
        except ValueError as error:
            result.append({
                'row': number, 'valid': False,
                'errors': [{'field': 'fx_rate', 'message': str(error)}],
            })
        else:
            result.append({
                'row': number, 'valid': True, 'data': normalized, 'errors': [],
                'instrument_resolution': 'resolved' if resolution else 'not_requested',
            })
    return {
        'digest': sha256(content).hexdigest(),
        'filename': Path(filename).name[:255],
        'rows': result,
        'valid': all(row['valid'] for row in result),
    }
