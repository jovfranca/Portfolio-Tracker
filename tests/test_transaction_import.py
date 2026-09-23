from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace as Obj

from openpyxl import Workbook
import pytest

from src.schemas import TransactionInput
from src.transaction_import import preview_import, read_rows


def csv_bytes(rows):
    header = (
        'ticker,broker,type,trade_date,settlement_date,quantity,'
        'unit_price,transaction_currency,fx_rate\n'
    )
    return (header + rows).encode()


def test_csv_preview_normalizes_rows_and_uses_settlement_fx(monkeypatch):
    calls = []

    def rates(session, currency, rate_type, reference_date):
        calls.append((currency, rate_type, reference_date))
        return [Obj(rate_side='MARKET', rate=Decimal('5.123456789012'))]

    monkeypatch.setattr('src.services.get_rates', rates)
    result = preview_import(
        object(),
        'transactions.csv',
        csv_bytes('aapl,Example,compra,2024-01-02,2024-01-04,2.5,100.01,usd,\n'),
    )
    assert result['valid']
    assert result['rows'][0]['data'] | {'fx_rate': 'ignored'} == {
        'trade_date': '2024-01-02',
        'settlement_date': '2024-01-04',
        'type': 'Buy',
        'asset': 'AAPL',
        'broker': 'Example',
        'allocation_class': 'Sem classe',
        'quantity': '2.5',
        'price': '100.01',
        'transaction_currency': 'USD',
        'fx_rate': 'ignored',
        'brokerage_fee': '0',
        'other_fees': '0',
        'notes': '',
    }
    assert result['rows'][0]['data']['fx_rate'] == '5.123456789012'
    assert calls == [('USD', 'FX', date(2024, 1, 4))]


def test_preview_reports_each_invalid_row_without_saving():
    result = preview_import(
        object(), 'bad.csv',
        csv_bytes('AAPL,Example,Buy,2024-01-02,2024-01-01,-2,100,USD,5\n'),
    )
    assert not result['valid']
    assert result['rows'][0]['row'] == 2
    assert {error['field'] for error in result['rows'][0]['errors']} >= {'quantity'}


@pytest.mark.parametrize('status', ['unresolved', 'ambiguous'])
def test_preview_explicitly_rejects_unresolved_or_ambiguous_instruments(monkeypatch, status):
    class EmptyResult:
        def all(self):
            return []

    class Session:
        def execute(self, _query):
            return EmptyResult()

    monkeypatch.setattr(
        'src.transaction_import.resolve_instrument',
        lambda *args, **kwargs: Obj(status=status, instrument=None),
    )
    result = preview_import(
        Session(), 'unknown.csv',
        csv_bytes('UNKNOWN,Example,Buy,2024-01-02,2024-01-03,1,10,BRL,1\n'),
        portfolio_id=1,
    )
    assert not result['valid']
    assert result['rows'][0]['instrument_resolution'] == status
    assert result['rows'][0]['errors'][0]['field'] == 'asset'


def test_xlsx_reader_accepts_required_columns():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['ticker', 'broker', 'type', 'trade_date', 'settlement_date', 'quantity', 'unit_price', 'transaction_currency'])
    sheet.append(['PETR4.SA', 'Example', 'Sell', date(2024, 1, 2), date(2024, 1, 3), 1, 30, 'BRL'])
    content = BytesIO()
    workbook.save(content)
    rows = read_rows('transactions.xlsx', content.getvalue())
    assert rows[0]['asset'] == 'PETR4.SA'
    assert rows[0]['trade_date'] == '2024-01-02'


def test_brl_rate_is_always_one():
    payload = TransactionInput.model_validate({
        'trade_date': '2024-01-02', 'settlement_date': '2024-01-03',
        'type': 'Buy', 'asset': 'PETR4.SA', 'broker': 'Example',
        'quantity': '0.1', 'price': '10.01', 'transaction_currency': 'brl', 'fx_rate': '9',
    })
    assert payload.fx_rate == Decimal('1')
    assert payload.quantity * payload.price == Decimal('1.001')


def test_resolved_brl_currency_overrides_supplied_fx_rate():
    from src.services import transaction_values
    payload = TransactionInput.model_validate({
        'trade_date': '2024-01-02', 'settlement_date': '2024-01-03',
        'type': 'Buy', 'asset': 'PETR4', 'broker': 'Example',
        'quantity': '1', 'price': '10', 'fx_rate': '9',
    })
    assert payload.transaction_currency is None
    assert transaction_values(None, payload, 'BRL')['fx_rate'] == Decimal('1')


@pytest.mark.parametrize('currency, rate', [('BRL', None), ('USD', '5')])
def test_known_fx_allows_pending_settlement(currency, rate):
    from datetime import timedelta
    payload = TransactionInput.model_validate({
        'trade_date': date.today(), 'settlement_date': date.today() + timedelta(days=2),
        'type': 'Buy', 'asset': 'TEST', 'broker': 'Example',
        'quantity': '1', 'price': '10', 'transaction_currency': currency, 'fx_rate': rate,
    })
    assert payload.fx_rate == Decimal(rate or '1')


@pytest.mark.parametrize('header', ['ticker,ticker', 'ticker,asset'])
def test_ambiguous_csv_headers_are_rejected(header):
    with pytest.raises(ValueError, match='coluna'):
        read_rows('bad.csv', (header + '\nAAPL,PETR4.SA\n').encode())


def test_unresolved_import_can_defer_currency_until_instrument_selection():
    result = preview_import(object(), 'bad.csv', csv_bytes(
        'AAPL,Example,Buy,2024-01-02,2024-01-03,1,100,,\n'))
    assert not result['valid']
    assert result['rows'][0]['errors'][0]['field'] == 'fx_rate'


def test_csv_extra_cells_are_not_silently_discarded():
    with pytest.raises(ValueError):
        read_rows('bad.csv', csv_bytes('AAPL,Example,Buy,2024-01-02,2024-01-03,1,100,USD,5,extra\n'))


def test_xlsx_duplicate_headers_are_rejected():
    workbook = Workbook()
    workbook.active.append(['ticker', 'asset'])
    workbook.active.append(['AAPL', 'PETR4.SA'])
    content = BytesIO()
    workbook.save(content)
    with pytest.raises(ValueError, match='coluna'):
        read_rows('bad.xlsx', content.getvalue())


def test_xlsx_expansion_limit(monkeypatch):
    workbook = Workbook()
    workbook.active.append(['ticker'])
    workbook.active.append(['AAPL'])
    content = BytesIO()
    workbook.save(content)
    monkeypatch.setattr('src.transaction_import.MAX_XLSX_EXPANDED_BYTES', 10, raising=False)
    with pytest.raises(ValueError):
        read_rows('bad.xlsx', content.getvalue())


@pytest.mark.parametrize('suffix', ['csv', 'xlsx'])
def test_preview_errors_use_source_line_numbers(suffix):
    content = csv_bytes('\nAAPL,Example,Buy,2024-01-02,2024-01-03,-1,100,BRL,1\n')
    if suffix == 'xlsx':
        workbook = Workbook()
        workbook.active.append(content.decode().splitlines()[0].split(','))
        workbook.active.append([])
        workbook.active.append(['AAPL', 'Example', 'Buy', '2024-01-02', '2024-01-03', -1, 100, 'BRL', 1])
        buffer = BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
    result = preview_import(object(), 'bad.' + suffix, content)
    assert not result['valid']
    assert result['rows'][0]['row'] == 3


def test_upload_stops_reading_as_soon_as_limit_is_exceeded(monkeypatch):
    import asyncio
    from fastapi import HTTPException
    from src.api.routes import import_preview
    monkeypatch.setattr('src.api.routes.get_portfolio', lambda *args: None)
    monkeypatch.setattr('src.api.routes.MAX_IMPORT_BYTES', 5)
    class Request:
        async def stream(self):
            yield b'1234'
            yield b'56'
            pytest.fail('Read beyond upload limit')
    with pytest.raises(HTTPException) as error:
        asyncio.run(import_preview(1, 'large.csv', Request(), object()))
    assert error.value.status_code == 413


def test_downloaded_template_matches_import_contract_without_network(monkeypatch):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from src.main import app
    from src.transaction_import import ALIASES

    monkeypatch.setattr('src.services.get_rates', lambda *args: pytest.fail('Template must not require FX lookup'))
    response = TestClient(app).get('/api/transactions/import-template.xlsx')
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'modelo-transacoes.xlsx' in response.headers['content-disposition']
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    headers = next(workbook.active.values)
    assert len(headers) == len(TransactionInput.model_fields)
    assert {ALIASES.get(name, name) for name in headers} == set(TransactionInput.model_fields)
    workbook.close()
    result = preview_import(object(), 'modelo.xlsx', response.content)
    assert result['valid'], result
    assert len(result['rows']) == 3
    assert {row['data']['type'] for row in result['rows']} == {'Buy', 'Sell'}
    assert {row['data']['transaction_currency'] for row in result['rows']} == {'BRL', 'USD'}
    assert all('fictício' in row['data']['notes'] for row in result['rows'])
